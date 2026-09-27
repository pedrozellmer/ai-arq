# -*- coding: utf-8 -*-
"""A planilha que o cliente recebe não fala "match".

🩸 27/09/2026 — A revisão do plano de outubro do Instagram (achado COPY-10) achou na aba
"Referências SINAPI" o cabeçalho "MATCH %" e a linha "↓ matches encontrados ↓" — no .xlsx do
exemplo público, que três posts de outubro mandam baixar, e em TODA planilha entregue, porque o
texto sai do gerador. Lendo o gerador inteiro, eram quatro: também "Sem match SINAPI" e o
rodapé "Como ler: MATCH %…". Jargão de dentro de casa; quem orça fala em correspondência,
candidato, semelhança.

O guarda gera uma planilha que passa pelos TRÊS caminhos da aba — item com escolhido +
alternativa, item com todos os candidatos reprovados pela IA, item sem candidato nenhum — e varre
TODAS as células de TODAS as abas.
"""
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from models import BudgetItem, Confidence, ProjectData  # noqa: E402

JARGAO = re.compile(r"match", re.I)

_BACIA = "Bacia sanitária com caixa acoplada"
_PENDENTE = "Pendente decorativo sobre bancada"
_CUBA = "Cuba de apoio + torneira (copa)"


def _itens():
    escolhido = {"codigo": "86888", "unidade": "UN", "similarity": 0.9, "_llm_picked": True,
                 "descricao": "BACIA SANITÁRIA COM CAIXA ACOPLADA LOUÇA BRANCA - FORNECIMENTO E INSTALAÇÃO"}
    alternativa = {"codigo": "100849", "unidade": "UN", "similarity": 0.4,
                   "descricao": "ASSENTO SANITÁRIO CONVENCIONAL - FORNECIMENTO E INSTALAÇÃO"}
    reprovado = {"codigo": "104567", "unidade": "UN", "similarity": 0.5, "_llm_rejected": True,
                 "descricao": "SPRINKLER TIPO PENDENTE"}
    return [
        BudgetItem(item_num="1.1", description=_BACIA, unit="un", quantity=3.0,
                   discipline="Instalações Hidráulicas", confidence=Confidence.ESTIMADO,
                   sinapi_matches=[escolhido, alternativa]),
        BudgetItem(item_num="2.1", description=_PENDENTE, unit="un", quantity=3.0,
                   discipline="Iluminação", confidence=Confidence.ESTIMADO,
                   sinapi_matches=[reprovado]),
        BudgetItem(item_num="3.1", description=_CUBA, unit="un", quantity=1.0,
                   discipline="Instalações Hidráulicas", confidence=Confidence.ESTIMADO),
    ]


def _celulas(tmp_path):
    """{aba: [texto de cada célula preenchida]} da planilha gerada."""
    from openpyxl import load_workbook
    from spreadsheet import generate_spreadsheet
    caminho = str(tmp_path / "planilha.xlsx")
    generate_spreadsheet(ProjectData(name="projeto de teste"), _itens(), caminho)
    return {aba.title: [str(c.value) for linha in aba.iter_rows() for c in linha if c.value is not None]
            for aba in load_workbook(caminho).worksheets}


def test_nenhuma_celula_da_planilha_fala_match(tmp_path):
    achados = ["%s: %r" % (aba, v[:80])
               for aba, valores in _celulas(tmp_path).items() for v in valores if JARGAO.search(v)]
    assert not achados, "a planilha do cliente voltou a falar 'match': " + " | ".join(achados)


def test_CONTROLE_a_aba_SINAPI_passou_pelos_tres_caminhos(tmp_path):
    """Sem isto, uma planilha sem a aba SINAPI — ou com um caminho só — passaria verde sem ter
    conferido o texto dos outros."""
    celulas = _celulas(tmp_path)
    assert "Referências SINAPI" in celulas, "a planilha saiu sem a aba SINAPI: %s" % list(celulas)
    aba = celulas["Referências SINAPI"]
    for descricao in (_BACIA, _PENDENTE, _CUBA):
        assert descricao in aba, "o item %r não passou pela aba SINAPI" % descricao
    assert "✓ IA conferiu" in aba, "o caminho do candidato escolhido pela IA não rodou"
    assert any(v.endswith("% (texto)") for v in aba), "o caminho da alternativa não rodou"


def test_CONTROLE_o_guarda_REPROVA_os_textos_antigos():
    """🧪 Prova que morde, sem tocar em arquivo nenhum: os quatro textos que estavam no ar."""
    for velho in ("MATCH %", "↓ matches encontrados ↓",
                  "Sem match SINAPI (descrição muito específica)",
                  "Como ler: MATCH % indica a similaridade"):
        assert JARGAO.search(velho), "o guarda não pega %r" % velho
