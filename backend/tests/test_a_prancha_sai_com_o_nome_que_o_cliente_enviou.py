# -*- coding: utf-8 -*-
"""A prancha aparece pro cliente com o nome que ELE enviou — não o do conversor.

🩸 18/09/2026. Quando o ODA recusa um .dwg, o plano B (libredwg) gera
"<nome>_libredwg.dxf" e o motor grava ESSE nome em `ref_sheet`. Medido no
banco antes de escrever qualquer linha:

  · 2.818 linhas de cliente (26,7%) com "_libredwg" no nome — 56 jobs, 43 contas
  · 1.682 delas MOSTRAVAM o sufixo na planilha (27 contas); as outras 1.136 só
    não mostravam porque o corte de 35 caracteres da coluna comia o fim
  · a tela do projeto mostra o nome INTEIRO — lá são as 43 contas
  · o cliente que se cadastrou hoje: 77 de 77 linhas assim

E um segundo dano, invisível: `_projeto_ja_enviado` compara os nomes ENVIADOS
("planta.dwg") com os gravados ("planta_libredwg.dxf"). Nunca casam. O aviso
"você já mandou esse caderno" era cego pra todo DWG convertido.

🔑 O CONSERTO É NA SAÍDA, NÃO NO BANCO. `ref_sheet` é também a CHAVE que o
botão "Ver desenho" manda de volta em `/api/sheet?ref=`. Trocar o que se grava
quebraria isso e não alcançaria as 2.818 linhas já gravadas. Uma regra
(`engine_rules.nome_que_o_cliente_enviou`), aplicada onde o nome CHEGA ao
cliente (planilha, tela, chat) e onde ele é COMPARADO (caderno repetido).

🧬 A doença da casa é a receita repetida. Por isso a tela NÃO tem cópia da
regra em JS: `/api/items` calcula `prancha_exibida` no backend e a tela só
mostra. O guarda de divergência abaixo cobra que planilha, API e tela digam o
MESMO nome pro mesmo `ref_sheet`.

🚫 O que NÃO cobre: um .dwg que o ODA converteu sem sufixo sai como "X.dxf",
indistinguível de um DXF enviado pelo cliente — 9 linhas, 1 conta, medido.
"""
import json
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import engine_rules as er  # noqa: E402
from _jsbancada import funcao_js, motor  # noqa: E402

_RAIZ = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
_PROJETO_HTML = os.path.join(_RAIZ, "projeto.html")

_SUJO = "planta baixa_libredwg.dxf"
_LIMPO = "planta baixa.dwg"


# ── a regra, pura ───────────────────────────────────────────────────────────

def test_o_sufixo_do_conversor_volta_a_ser_o_dwg_do_cliente():
    assert er.nome_que_o_cliente_enviou(_SUJO) == _LIMPO
    # com o hint da IA depois — a forma "arquivo (hint)" que o motor grava
    assert er.nome_que_o_cliente_enviou(_SUJO + " (Térreo)") == _LIMPO + " (Térreo)"
    # parêntese que é parte do NOME (caso real de 17/09: "prancha (1)")
    assert er.nome_que_o_cliente_enviou("prancha (1)_libredwg.dxf") == "prancha (1).dwg"
    # caixa: o conversor não muda o radical; a extensão volta minúscula
    assert er.nome_que_o_cliente_enviou("PLANTA_LIBREDWG.DXF") == "PLANTA.dwg"


def test_CONTROLE_nome_sem_sufixo_passa_intocado():
    """🪤 Regra que 'limpa' demais some com nome legítimo. Um DXF enviado pelo
    cliente, um PDF, e um nome com "_libredwg" no MEIO ficam como estão."""
    for nome in ("planta.dxf", "planta.pdf (hint)", "x_libredwg_v2.dxf",
                 "relatorio_libredwg.pdf", "planta.dwg"):
        assert er.nome_que_o_cliente_enviou(nome) == nome, nome
    assert er.nome_que_o_cliente_enviou("") == ""
    assert er.nome_que_o_cliente_enviou(None) == ""


# ── a planilha REAL ─────────────────────────────────────────────────────────

def _planilha_com(ref_sheet):
    from models import BudgetItem, Confidence, ProjectData
    from spreadsheet import generate_spreadsheet
    from openpyxl import load_workbook
    it = BudgetItem(item_num="1.1", description="Alvenaria de vedação",
                    unit="m2", quantity=12.0, ref_sheet=ref_sheet,
                    confidence=Confidence.ESTIMADO, discipline="Paredes")
    out = os.path.join(tempfile.mkdtemp(prefix="prancha_nome_"), "p.xlsx")
    generate_spreadsheet(ProjectData(name="projeto de teste"), [it], out)
    wb = load_workbook(out)
    # 🪤 A 1ª versão procurava OBSERVAÇÕES — eu tinha lido errado. A
    # referência da prancha mora na coluna REF. (spreadsheet.py ~595/653), na
    # aba "Orçamento" (a ativa pode ser a capa), abaixo das notas de rodapé.
    ws = wb["Orçamento"] if "Orçamento" in wb.sheetnames else wb.active
    col = None
    for row in ws.iter_rows(min_row=1, max_row=40):
        for c in row:
            if str(c.value or "").strip().upper().rstrip(".") == "REF":
                col = c.column
                break
        if col:
            break
    assert col, "não achei a coluna REF. no cabeçalho da planilha"
    textos = [str(ws.cell(row=r, column=col).value or "") for r in range(1, ws.max_row + 1)]
    return "\n".join(t for t in textos if "Alvenaria" in t or ".dwg" in t or ".dxf" in t or ".pdf" in t)


def test_a_planilha_mostra_o_dwg_que_o_cliente_enviou():
    obs = _planilha_com(_SUJO)
    assert "_libredwg" not in obs, obs
    assert _LIMPO in obs, "a coluna não traz o nome que o cliente enviou: %r" % obs


def test_CONTROLE_a_planilha_nao_mexe_em_nome_de_PDF():
    obs = _planilha_com("planta baixa.pdf")
    assert "planta baixa.pdf" in obs, obs
    assert ".dwg" not in obs


# ── o aviso de caderno repetido, EXECUTADO ──────────────────────────────────

def _ja_enviado(monkeypatch, gravados, novos):
    import main
    chamadas = []

    def _rows(metodo, caminho, **k):
        chamadas.append(caminho)
        if caminho.startswith("/projects?"):
            return [{"job_id": "job-velho", "project_name": "caderno",
                     "created_at": "2026-09-01T12:00:00Z",
                     "user_pe_direito": None, "user_total_area": None}]
        if caminho.startswith("/project_items?"):
            return [{"ref_sheet": g} for g in gravados]
        return []
    monkeypatch.setattr(main, "_supa_rows", _rows)
    # 🪤 22/09/2026: Storage mudo — estes guardas medem a reserva pelos itens,
    # e sem isto a função sairia pra rede de verdade.
    monkeypatch.setattr(main, "_pranchas_com_tamanho", lambda *a, **k: None)
    r = main._projeto_ja_enviado("user-1", set(novos), None, None)
    assert chamadas, "a função não consultou nada — o guarda não executou o caminho"
    return r


def test_o_aviso_de_caderno_repetido_enxerga_o_DWG_convertido(monkeypatch):
    """🩸 O cliente que reenvia é justamente quem achou que o problema era o
    arquivo dele. Com o sufixo no gravado, 3 iguais viravam 0 iguais."""
    r = _ja_enviado(monkeypatch,
                    gravados=["a_libredwg.dxf", "b_libredwg.dxf", "c_libredwg.dxf"],
                    novos=["a.dwg", "b.dwg", "c.dwg"])
    assert r and r["n_iguais"] == 3, r


def test_CONTROLE_caderno_diferente_continua_sem_aviso(monkeypatch):
    r = _ja_enviado(monkeypatch,
                    gravados=["a_libredwg.dxf", "b_libredwg.dxf", "c_libredwg.dxf"],
                    novos=["x.dwg", "y.dwg", "z.dwg"])
    assert r is None, r


def test_o_aviso_enxerga_acento_e_extensao_trocada_pelo_conversor(monkeypatch):
    """🩸 Achado da revisão: a cegueira era de TODO DWG, não só do libredwg. O
    gravado nasce sanitizado (sem acento) e o ODA troca .dwg por .dxf; o
    enviado chega cru do navegador. Três razões pra nunca casar."""
    r = _ja_enviado(monkeypatch,
                    gravados=["galpao.dxf", "corte a-a.dxf", "fachada_libredwg.dxf"],
                    novos=["Galpão.dwg", "CORTE A-A.dwg", "fachada.DWG"])
    assert r and r["n_iguais"] == 3, r


def test_CONTROLE_mandar_o_DWG_do_PDF_que_ja_enviou_NAO_e_repeticao(monkeypatch):
    """🚫 De propósito: .pdf e .dwg do mesmo nome são desenhos diferentes pra
    medição — quem manda o DWG depois do PDF está fazendo o que o aviso pede.
    Reduzir tudo ao radical acusaria justamente esse cliente."""
    r = _ja_enviado(monkeypatch,
                    gravados=["planta.pdf (Térreo)", "corte.pdf", "fachada.pdf"],
                    novos=["planta.dwg", "corte.dwg", "fachada.dwg"])
    assert r is None, r


def test_CONTROLE_o_mesmo_PDF_de_novo_continua_avisando(monkeypatch):
    r = _ja_enviado(monkeypatch,
                    gravados=["planta.pdf (Térreo)", "corte.pdf", "fachada.pdf"],
                    novos=["planta.pdf", "corte.pdf", "fachada.pdf"])
    assert r and r["n_iguais"] == 3, r


# ── a API que alimenta a tela ───────────────────────────────────────────────

def _itens_da_api(monkeypatch, ref_sheet):
    import main
    import urllib.request
    monkeypatch.setattr(main, "_require_project_owner", lambda *a, **k: None)
    monkeypatch.setattr(main, "_itens_do_projeto_completos",
                        lambda job: ([{"id": 1, "ref_sheet": ref_sheet,
                                       "description": "x"}], True))

    def _sem_rede(*a, **k):
        raise RuntimeError("sem rede no teste")
    monkeypatch.setattr(urllib.request, "urlopen", _sem_rede)
    return main.get_project_items("job-teste", object())["items"]


def test_a_API_entrega_o_nome_limpo_SEM_mexer_na_chave(monkeypatch):
    """🔑 Duas coisas no mesmo item: o rótulo limpo e a chave intacta. Se
    `ref_sheet` mudasse, o botão "Ver desenho" mandaria um nome que
    `/api/sheet` não conhece."""
    it = _itens_da_api(monkeypatch, _SUJO)[0]
    assert it["prancha_exibida"] == _LIMPO, it
    assert it["ref_sheet"] == _SUJO, "a CHAVE foi alterada — quebra o Ver desenho"


# ── a tela, RODANDO o JavaScript real ───────────────────────────────────────

_MAP_POLYFILL = """
function Map(){ this._k=[]; this._v=[]; }
Map.prototype.get=function(k){ var i=this._k.indexOf(k); return i<0?undefined:this._v[i]; };
Map.prototype.set=function(k,v){ var i=this._k.indexOf(k); if(i<0){this._k.push(k);this._v.push(v);} else {this._v[i]=v;} return this; };
Map.prototype.has=function(k){ return this._k.indexOf(k)>=0; };
Object.defineProperty(Map.prototype,'size',{get:function(){ return this._k.length; }});
Map.prototype.entries=function(){ var o=[]; for(var i=0;i<this._k.length;i++) o.push([this._k[i],this._v[i]]); return o; };
true;
"""
# 🪤 O `true;` no fim NÃO é enfeite: o valor do último comando do prelúdio é o
# que `evaljs` devolve, e dukpy não consegue devolver uma FUNÇÃO — o polyfill
# terminando em `Map.prototype.entries = function…` derrubava o motor com
# "Invalid Result Value" antes de qualquer teste rodar.


def _mapa_da_tela(itens):
    js = motor(_MAP_POLYFILL)
    js.evaljs(funcao_js("collectCadFiles", _PROJETO_HTML))
    js.evaljs("var ITENS = %s;" % json.dumps(itens))
    bruto = js.evaljs("JSON.stringify(collectCadFiles(ITENS).entries())")
    return {k: v for k, v in json.loads(bruto)}


def test_a_tela_mostra_o_nome_limpo_e_guarda_a_chave_crua():
    m = _mapa_da_tela([{"ref_sheet": _SUJO, "prancha_exibida": _LIMPO,
                        "confidence": "estimado", "quantity": 1}])
    assert list(m.keys()) == [_SUJO], "a chave do mapa mudou — o botão perde a prancha"
    assert m[_SUJO]["exibida"] == _LIMPO, m


def test_CONTROLE_sem_prancha_exibida_a_tela_cai_no_nome_cru():
    """🪤 Resposta velha da API (sem o campo) não pode deixar o rótulo vazio."""
    m = _mapa_da_tela([{"ref_sheet": "planta.pdf (hint da IA)",
                        "confidence": "estimado", "quantity": 1}])
    assert m["planta.pdf"]["exibida"] == "planta.pdf", m


# ── as três superfícies dizem o MESMO nome ──────────────────────────────────

def test_planilha_API_e_tela_concordam(monkeypatch):
    """🧬 O guarda mira na DIVERGÊNCIA entre as cópias — é assim que a receita
    repetida deixa de ser doença. Mesmo `ref_sheet`, mesmo nome nas três."""
    api = _itens_da_api(monkeypatch, _SUJO)[0]["prancha_exibida"]
    tela = _mapa_da_tela([{"ref_sheet": _SUJO, "prancha_exibida": api,
                           "confidence": "estimado", "quantity": 1}])[_SUJO]["exibida"]
    planilha = _planilha_com(_SUJO)
    assert api == tela == _LIMPO, (api, tela)
    assert api in planilha, planilha


def test_o_chat_le_a_planilha_com_o_nome_do_cliente():
    import agent
    assert agent._nome_da_prancha_para_o_cliente(_SUJO) == _LIMPO
    assert agent._nome_da_prancha_para_o_cliente(None) == ""


def test_o_chat_le_a_COLUNA_certa_da_planilha():
    """🩸 Achado da revisão: o chat lia até a 9ª coluna e chamava a 9ª de
    `ref_sheet` — mas a 9ª é ORIGEM DA MEDIÇÃO; a referência é a 11ª, REF.
    O chat entregava a origem da medição como se fosse o nome da prancha.
    Este guarda gera a planilha REAL e lê como o chat lê."""
    import agent
    from models import BudgetItem, Confidence, ProjectData
    from spreadsheet import generate_spreadsheet
    from openpyxl import load_workbook
    it = BudgetItem(item_num="1.1", description="Alvenaria de vedação",
                    unit="m2", quantity=12.0, ref_sheet=_SUJO,
                    confidence=Confidence.CONFIRMADO, discipline="Paredes",
                    origem="dxf_geom")
    out = os.path.join(tempfile.mkdtemp(prefix="chat_coluna_"), "p.xlsx")
    generate_spreadsheet(ProjectData(name="projeto de teste"), [it], out)
    linhas = list(agent._iter_orcamento_rows(load_workbook(out)))
    # 🪤 A planilha real abre com linhas de METADADO (0.1, 0.9…) antes do
    # item; a 1ª versão deste guarda pegava `linhas[0]` e lia o ref vazio da
    # linha 0.1 — reprovando o código certo. Pega a linha do ITEM.
    do_item = [l for l in linhas if l["item_num"] == "1.1"]
    assert do_item, "o chat não achou a linha 1.1 na aba Orçamento: %r" % [
        l["item_num"] for l in linhas]
    ref = do_item[0]["ref_sheet"]
    assert ref == _LIMPO, (
        "o chat leu %r como nome da prancha — coluna errada ou sem limpeza" % ref)
    # controle: o que o chat chama de ref NÃO é a origem da medição
    assert "geometria" not in ref.lower() and "medid" not in ref.lower(), ref
