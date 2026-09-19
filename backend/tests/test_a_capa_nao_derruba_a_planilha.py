# -*- coding: utf-8 -*-
"""Um campo decorativo da capa não pode derrubar a entrega inteira.

🩸 19/09/2026, cliente real, de manhã (job 92f27fbb). O motor leu duas
pranchas, converteu o DWG pelo caminho reserva, juntou 139 itens com 16
medidos — 25 minutos de máquina — e, ao escrever a CAPA da planilha, estourou:

    AttributeError: 'str' object has no attribute 'get'
    spreadsheet.py:342 -> name = dept.get('name', '')

O job virou `status=error`, os 139 itens NUNCA foram gravados (ficaram os 76
da primeira passada, do primeiro arquivo), e o cliente recebeu um e-mail
pedindo pra TROCAR O ARQUIVO — por um defeito que era nosso, num desenho que
o motor tinha acabado de ler inteiro.

A causa: `models.py` declara `departments: list[dict]`, o prompt pede objetos
com `name`/`positions`, e `engine_rules.mesclar_project_data` copiava o que a
IA mandasse. Naquele job a IA devolveu uma lista de STRINGS. **Tipo declarado
não é contrato enquanto ninguém o aplica na porta de entrada.**

Dois guardas, porque o conserto tem duas camadas e cada uma sozinha deixa
passar um jeito de quebrar:
  1. a ENTRADA normaliza (`departamentos_no_formato_do_modelo`) — assim o resto
     do sistema pode confiar no tipo;
  2. a ESCRITA tolera — porque quem escreve a entrega não pode cair por causa
     de um cabeçalho, venha o dado por onde vier (inclusive de um caminho que
     ainda não existe hoje). O vizinho `new_rooms` já fazia isso desde sempre.
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import engine_rules  # noqa: E402
from models import ProjectData  # noqa: E402


# ─────────────────────────────────────────────────────────────────────────────
#  (1) A porta de entrada: o que a IA manda sai no formato do modelo
# ─────────────────────────────────────────────────────────────────────────────

def test_lista_de_STRINGS_da_IA_vira_o_formato_do_modelo():
    """O formato exato que quebrou a planilha do cliente."""
    saida = engine_rules.departamentos_no_formato_do_modelo(
        ["Recepção", "Sala de reunião", "Copa"])
    assert saida == [{"name": "Recepção", "positions": 0},
                     {"name": "Sala de reunião", "positions": 0},
                     {"name": "Copa", "positions": 0}], saida
    assert all(isinstance(d, dict) for d in saida)


def test_o_formato_certo_atravessa_intacto():
    saida = engine_rules.departamentos_no_formato_do_modelo(
        [{"name": "Diretoria", "positions": 4}])
    assert saida == [{"name": "Diretoria", "positions": 4}], saida


def test_mistura_e_lixo_nao_viram_linha_vazia_na_capa():
    """🪤 Antes, um item ilegível viraria a linha "  None" impressa na capa da
    planilha do cliente. Descarta-se em silêncio o que não dá pra entender."""
    saida = engine_rules.departamentos_no_formato_do_modelo(
        ["Recepção", {"name": "TI", "positions": "12"}, None, "", {"positions": 3},
         {"nome": "Almoxarifado", "posicoes": 2}])
    assert saida == [{"name": "Recepção", "positions": 0},
                     {"name": "TI", "positions": 12},
                     {"name": "Almoxarifado", "positions": 2}], saida


def test_um_texto_solto_tambem_e_aceito():
    """A IA já devolveu campo de lista como string crua em outros pontos."""
    assert engine_rules.departamentos_no_formato_do_modelo("Recepção") == [
        {"name": "Recepção", "positions": 0}]


def test_CONTROLE_a_mesclagem_USA_a_normalizacao(monkeypatch):
    """Guarda de prato: a função pode estar certa e não ser chamada — foi
    exatamente assim que este defeito existiu."""
    destino = ProjectData()
    engine_rules.mesclar_project_data(destino, {"departments": ["Recepção"]})
    assert destino.departments == [{"name": "Recepção", "positions": 0}], (
        "a mesclagem voltou a copiar o que a IA mandou, sem passar pela porta")


# ─────────────────────────────────────────────────────────────────────────────
#  (2) A escrita: a entrega não cai por causa da capa
# ─────────────────────────────────────────────────────────────────────────────

def _planilha_com(departments, tmp_path, typology="office"):
    from models import BudgetItem, Confidence
    from spreadsheet import generate_spreadsheet
    projeto = ProjectData(name="projeto de teste")
    projeto.departments = departments
    itens = [BudgetItem(item_num="1.1", description="Alvenaria de vedação",
                        unit="m2", quantity=12.0, discipline="Paredes",
                        confidence=Confidence.ESTIMADO)]
    caminho = str(tmp_path / "saida.xlsx")
    generate_spreadsheet(projeto, itens, caminho, typology=typology)
    return caminho


def _texto_da_capa(caminho):
    from openpyxl import load_workbook
    aba = load_workbook(caminho).worksheets[0]
    return "\n".join(str(c.value or "") for linha in aba.iter_rows() for c in linha)


def test_a_planilha_SAI_mesmo_com_o_formato_que_quebrou(tmp_path):
    """🚨 O caso do cliente, ponta a ponta: com a lista de strings, a planilha
    tem que ser escrita do mesmo jeito — 139 itens não se perdem por uma capa."""
    caminho = _planilha_com(["Recepção", "Copa"], tmp_path)
    assert os.path.getsize(caminho) > 0
    texto = _texto_da_capa(caminho)
    assert "Recepção" in texto, "o departamento não chegou à capa: %s" % texto[:400]
    assert "None" not in texto, "a capa imprimiu None: %s" % texto[:400]


def test_a_planilha_SAI_com_lixo_no_lugar_do_departamento(tmp_path):
    """Nenhum formato de entrada derruba a escrita — nem os que ainda não
    aconteceram."""
    for formato in (["Recepção", None, 42, {"name": "TI", "positions": 3}],
                    [{"sem_nome": 1}],
                    ["", "   "]):
        caminho = _planilha_com(formato, tmp_path)
        assert os.path.getsize(caminho) > 0, formato


def test_CONTROLE_a_capa_continua_escrevendo_o_departamento_certo(tmp_path):
    """O outro lado: tolerar não pode virar ignorar — o dado bom continua saindo,
    com as posições."""
    caminho = _planilha_com([{"name": "Diretoria", "positions": 4}], tmp_path)
    texto = _texto_da_capa(caminho)
    assert "Diretoria" in texto and "4 posições" in texto, texto[:400]
    assert "DEPARTAMENTOS" in texto, "o cabeçalho da seção sumiu: %s" % texto[:400]
