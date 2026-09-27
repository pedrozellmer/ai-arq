# -*- coding: utf-8 -*-
"""Gera o exemplo público — a planilha E a tabela da página — a partir de UMA fonte: exemplos/itens.json.

🩸 27/09/2026 — A planilha do exemplo era de 17/07 e o gerador seguiu andando: em 24/08 a planilha
ganhou ORIGEM DA MEDIÇÃO e ESPECIFICAÇÃO (9 → 11 colunas); em 27/09 a aba SINAPI parou de falar
"match" e a nota de perdas mudou. Quem baixava o exemplo pra "ver como volta" recebia uma planilha
que não sai mais (achado PROD-11 da revisão do plano de outubro do Instagram — o exemplo é o 1º link
da bio o mês inteiro). Página e planilha já tinham se descolado uma vez (d5b45d6, 17/07): por isso
as duas saem da MESMA rodada, e a numeração da página é LIDA da planilha gerada.

Uso:
    python scripts/gerar_exemplo_publico.py             # regera o .xlsx e a tabela do exemplo.html
    python scripts/gerar_exemplo_publico.py --conferir  # só confere; sai 1 se algo ficou velho

Nada aqui chama IA nem busca no SINAPI: os candidatos da rodada de 17/07 estão gravados no itens.json.
"""
import argparse
import io
import json
import math
import os
import re
import sys
import tempfile

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(RAIZ, "backend"))

FONTE = os.path.join(RAIZ, "exemplos", "itens.json")
XLSX = os.path.join(RAIZ, "exemplos", "quantitativo-exemplo-aiarq.xlsx")
PAGINA = os.path.join(RAIZ, "exemplo.html")
_BLOCO_EX = re.compile(r"const EX = \{.*?\]\};", re.S)
_NUM_ITEM = re.compile(r"^\d+\.\d+$")
_NUM_INDIRETO = re.compile(r"^S\.\d+$")


def carregar():
    with io.open(FONTE, encoding="utf-8") as f:
        return json.load(f)


def _itens_do_modelo(dados, numeros=None):
    """BudgetItems do jeito que o motor entrega pro gerador. `numeros` (desc → nº) vem da 1ª
    passada: a aba SINAPI escreve o item_num que recebe, e a aba Orçamento renumera por
    disciplina — sem a 2ª passada as duas abas dariam números diferentes pro mesmo item."""
    from models import BudgetItem, Confidence
    from spreadsheet import REF_MIN_CONFIDENCE
    itens = []
    fila = dados["itens"]
    if numeros:
        # 2ª passada na ordem da numeração: a aba SINAPI lista na ordem em que recebe, e assim
        # ela lê 1.1, 1.2, 2.1… como a aba Orçamento (dentro da disciplina, a ordem não muda).
        fila = sorted(fila, key=lambda i: _chave_num(numeros[i["desc"]]))
    for it in fila:
        matches = []
        for c in it["sinapi"]:
            m = {"codigo": c["codigo"], "descricao": c["descricao"], "unidade": c["unidade"]}
            if c.get("busca_simplificada"):
                m["_match_level"] = "simplified"
            if c["papel"] == "escolhido":
                # A nota de texto do escolhido não aparece na planilha ("✓ IA conferiu"); só
                # precisa passar do corte da coluna REF, como passou na rodada de 17/07.
                m["_llm_picked"] = True
                m["similarity"] = REF_MIN_CONFIDENCE
            elif c["papel"] == "reprovado":
                m["_llm_rejected"] = True
            else:
                m["similarity"] = c["semelhanca_texto"]
            matches.append(m)
        if matches and any(c["papel"] == "escolhido" for c in it["sinapi"]):
            assert it["sinapi"][0]["papel"] == "escolhido", (
                "%r: o escolhido tem que vir primeiro — é ele que vai pra coluna REF" % it["desc"])
        itens.append(BudgetItem(
            item_num=(numeros or {}).get(it["desc"], "0"),
            description=it["desc"], unit=it["un"], quantity=it["qtd"], discipline=it["disc"],
            observations=it["obs"],
            confidence=Confidence.CONFIRMADO if it["medido"] else Confidence.ESTIMADO,
            origem="dxf_geom" if it["medido"] else "",
            sinapi_matches=matches))
    return itens


def _projeto(dados):
    from models import ProjectData
    p = dados["projeto"]
    return ProjectData(name=p["name"], address=p["address"], architect=p["architect"], phase=p["phase"],
                       total_area=p["total_area"], layout_area=p["layout_area"],
                       no_intervention_area=p["no_intervention_area"])


def _ler_orcamento(caminho, descricoes):
    """(desc → nº na aba Orçamento, quantas linhas de custo indireto S.x)."""
    from openpyxl import load_workbook
    ws = load_workbook(caminho)["Orçamento"]
    numeros, indiretos = {}, 0
    for linha in ws.iter_rows(values_only=True):
        num, desc = (str(linha[0]) if linha[0] is not None else ""), linha[1]
        if _NUM_ITEM.match(num) and desc in descricoes:
            numeros[desc] = num
        elif _NUM_INDIRETO.match(num):
            indiretos += 1
    return numeros, indiretos


def gerar_planilha(dados, destino):
    """Roda o gerador de verdade (2 passadas) e devolve (desc → nº, nº de linhas de custo indireto)."""
    from spreadsheet import generate_spreadsheet
    descricoes = {it["desc"] for it in dados["itens"]}
    tipologia = dados["projeto"]["tipologia"]
    with tempfile.TemporaryDirectory() as tmp:
        rascunho = os.path.join(tmp, "passada1.xlsx")
        generate_spreadsheet(_projeto(dados), _itens_do_modelo(dados), rascunho, typology=tipologia)
        numeros, _ = _ler_orcamento(rascunho, descricoes)
    faltando = descricoes - set(numeros)
    assert not faltando, "o gerador não escreveu estes itens na aba Orçamento: %s" % sorted(faltando)
    generate_spreadsheet(_projeto(dados), _itens_do_modelo(dados, numeros), destino, typology=tipologia)
    numeros2, indiretos = _ler_orcamento(destino, descricoes)
    assert numeros2 == numeros, "a numeração mudou entre as duas passadas"
    return numeros, indiretos


def _chave_num(num):
    a, b = num.split(".")
    return int(a), int(b)


def ex_da_pagina(dados, numeros):
    """O objeto EX da página, na ordem da planilha. A composição vai cortada em 70 caracteres —
    a página emenda "…" e o código do lado; o texto inteiro está na aba SINAPI.

    27/09/2026 (Pedro liberou): o exemplo mostrava 72% medido, 8 de 11 áreas seladas e nenhuma
    linha em branco — na base real, área sai medida em 1,8% das linhas e 31,6% das linhas em CAD
    voltam sem quantidade. Agora há três estados, como na planilha: medido, estimado com número e
    em branco (estimado com quantidade zero, que a planilha deixa vazia pra você preencher). E a
    observação vai junto: é nela que está a conta de cada estimativa — ou a falta de origem."""
    linhas = []
    for it in sorted(dados["itens"], key=lambda i: _chave_num(numeros[i["desc"]])):
        esc = next((c for c in it["sinapi"] if c["papel"] == "escolhido"), None)
        linhas.append({"num": numeros[it["desc"]], "desc": it["desc"], "un": it["un"], "qtd": it["qtd"],
                       "disc": it["disc"], "medido": it["medido"], "obs": it["obs"],
                       "sinapi": ({"cod": esc["codigo"], "comp": esc["descricao"][:70], "un": esc["unidade"]}
                                  if esc else None)})
    medidos = sum(1 for i in dados["itens"] if i["medido"])
    brancos = sum(1 for i in dados["itens"] if not i["medido"] and not i["qtd"])
    return {"total": len(linhas), "medidos": medidos, "estimados": len(linhas) - medidos - brancos,
            "brancos": brancos, "area": dados["projeto"]["total_area"], "itens": linhas}


def percentuais(ex):
    """(% medido, % estimado, % em branco), somando 100 — as MESMAS contas do JS da página.
    Arredonda como o Math.round do JS (meio pra cima), não como o round do Python (meio pro par)."""
    pm = int(math.floor(100.0 * ex["medidos"] / ex["total"] + 0.5))
    pb = int(math.floor(100.0 * ex["brancos"] / ex["total"] + 0.5))
    return pm, 100 - pm - pb, pb


def bloco_ex(ex, eol="\n"):
    """O texto do `const EX = {...};` no formato que a página já usava (1 item por linha)."""
    cab = json.dumps({k: ex[k] for k in ("total", "medidos", "estimados", "brancos", "area")},
                     ensure_ascii=False, separators=(",", ":"))[:-1]
    corpo = ("," + eol).join("  " + json.dumps(i, ensure_ascii=False) for i in ex["itens"])
    return "const EX = " + cab + ',"itens":[' + eol + corpo + eol + "  ]};"


# 27/09/2026 (auditoria de aquisição, C10): o <div id="itens"> chegava VAZIO no HTML
# servido — os 25 itens só existiam dentro do <script>. Robô que não roda JS (a maioria
# dos leitores de IA) via a página do exemplo sem o exemplo. Agora o gerador escreve uma
# tabela simples ali dentro, da MESMA fonte do `const EX`; o JS continua trocando pela
# versão visual. Sem <div> dentro, pra o fim do bloco não ser ambíguo.
_BLOCO_ITENS = re.compile(r'(<div id="itens" class="mt-4">)(.*?)(</div>)', re.S)


def _fmt_qtd(q):
    """Como o `fmtQt` do JS: pt-BR, até 2 casas, sem zero sobrando (118.5 → 118,5)."""
    inteiro, _, frac = ("%.2f" % q).partition(".")
    frac = frac.rstrip("0")
    inteiro = "{:,}".format(int(inteiro)).replace(",", ".")
    return inteiro + ("," + frac if frac else "")


def tabela_estatica(ex, eol="\n"):
    """A tabela dos itens em HTML puro, legível sem JavaScript."""
    import html as _h
    e = lambda s: _h.escape(str(s), quote=False)  # noqa: E731
    linhas = ['<table class="w-full text-sm">',
              "<caption>Os %d itens do exemplo: %d medidos da geometria, %d estimados e %d em branco</caption>"
              % (ex["total"], ex["medidos"], ex["estimados"], ex["brancos"]),
              "<thead><tr><th>Item</th><th>Descrição</th><th>Un.</th><th>Quantidade</th>"
              "<th>Confiança</th><th>SINAPI</th></tr></thead>", "<tbody>"]
    ordem, grupos = [], {}
    for it in ex["itens"]:
        if it["disc"] not in grupos:
            ordem.append(it["disc"])
            grupos[it["disc"]] = []
        grupos[it["disc"]].append(it)
    for disc in ordem:
        linhas.append('<tr><th colspan="6">%s</th></tr>' % e(disc))
        for it in grupos[disc]:
            sin = ("%s — %s… [%s]" % (e(it["sinapi"]["cod"]), e(it["sinapi"]["comp"]), e(it["sinapi"]["un"]))
                   if it["sinapi"] else "sem equivalente no SINAPI")
            # a observação vai na célula da descrição: é ela que diz de onde veio o número
            obs = e(it["obs"]) if it["obs"] else "sem origem escrita: confira de onde veio o número"
            linhas.append("<tr><td>%s</td><td>%s<br><small>%s</small></td><td>%s</td><td>%s</td><td>%s</td><td>%s</td></tr>"
                          % (e(it["num"]), e(it["desc"]), obs, e(it["un"]),
                             _fmt_qtd(it["qtd"]) if it["qtd"] else "em branco",
                             "Medido" if it["medido"] else "Estimado", sin))
    linhas += ["</tbody>", "</table>"]
    return eol.join(linhas)


def frases_que_a_pagina_tem_que_ter(ex, indiretos):
    """A prosa da página cita as contagens. Se o dado mudar, a prosa muda junto — ou o script para."""
    pm, pe, pb = percentuais(ex)
    n_disc = len({i["disc"] for i in ex["itens"]})
    n_sinapi = sum(1 for i in ex["itens"] if i["sinapi"])
    area = ("%.1f" % ex["area"]).replace(".", ",")
    return [
        "%d itens em %d disciplinas" % (ex["total"], n_disc),
        "%d medidos do desenho" % ex["medidos"],
        "%d estimados com a conta escrita" % ex["estimados"],
        "%d em branco" % ex["brancos"],
        "%d deles com refer" % n_sinapi,
        "os outros %d n" % (ex["total"] - n_sinapi),
        "%d linhas de custo indireto" % indiretos,
        "do que os %d itens abaixo" % ex["total"],
        'id="rx-med">%d<' % ex["medidos"],
        'id="rx-est">%d<' % ex["estimados"],
        'id="rx-bra">%d<' % ex["brancos"],
        'id="rx-bar-med" class="h-3 bg-emerald-500" style="width:%d%%"' % pm,
        'id="rx-bar-est" class="h-3 bg-amber-400" style="width:%d%%"' % pe,
        "%d%% medido · %d%% estimado · %d%% em branco" % (pm, pe, pb),
        "Neste exemplo, %d%%" % pm,
        "%s m" % area,
    ]


def _celulas(caminho):
    from openpyxl import load_workbook
    wb = load_workbook(caminho)
    return {ws.title: {c.coordinate: c.value for linha in ws.iter_rows() for c in linha if c.value is not None}
            for ws in wb.worksheets}


def diferencas_da_planilha(esperada, atual):
    """Compara VALOR de célula (o .xlsx carrega data de criação — bytes nunca batem)."""
    a, b = _celulas(esperada), _celulas(atual)
    difs = []
    if list(a) != list(b):
        difs.append("abas: %s × %s" % (list(a), list(b)))
    for aba in a:
        for coord in sorted(set(a[aba]) | set(b.get(aba, {}))):
            va, vb = a[aba].get(coord), b.get(aba, {}).get(coord)
            if va != vb:
                difs.append("%s!%s: gerador diz %r, arquivo tem %r" % (aba, coord, str(va)[:60], str(vb)[:60]))
    return difs


def conferir(pagina=PAGINA, xlsx=XLSX):
    """Lista do que está velho (vazia = tudo em dia). Não escreve nada."""
    dados = carregar()
    problemas = []
    with tempfile.TemporaryDirectory() as tmp:
        novo = os.path.join(tmp, "exemplo.xlsx")
        numeros, indiretos = gerar_planilha(dados, novo)
        problemas += ["planilha: " + d for d in diferencas_da_planilha(novo, xlsx)[:15]]
    ex = ex_da_pagina(dados, numeros)
    with io.open(pagina, encoding="utf-8", newline="") as f:
        html = f.read()
    m = _BLOCO_EX.search(html)
    if not m:
        problemas.append("página: não achei o bloco `const EX = {...};`")
    elif m.group(0).replace("\r\n", "\n") != bloco_ex(ex):
        problemas.append("página: a tabela (const EX) não é a que sai do itens.json + gerador")
    t = _BLOCO_ITENS.search(html)
    if not t:
        problemas.append('página: não achei o <div id="itens" class="mt-4">')
    elif t.group(2).replace("\r\n", "\n") != tabela_estatica(ex):
        problemas.append("página: a tabela estática (sem JS) não é a que sai do itens.json + gerador")
    problemas +=["página: a prosa não diz %r" % fr for fr in frases_que_a_pagina_tem_que_ter(ex, indiretos)
                  if fr not in html]
    return problemas


def regerar():
    dados = carregar()
    tmp_xlsx = XLSX + ".tmp.xlsx"
    numeros, indiretos = gerar_planilha(dados, tmp_xlsx)
    ex = ex_da_pagina(dados, numeros)
    with io.open(PAGINA, encoding="utf-8", newline="") as f:
        html = f.read()
    eol = "\r\n" if "\r\n" in html else "\n"
    assert _BLOCO_EX.search(html), "não achei o bloco `const EX = {...};` no exemplo.html"
    novo_html = _BLOCO_EX.sub(lambda _m: bloco_ex(ex, eol), html, count=1)
    assert _BLOCO_ITENS.search(novo_html), 'não achei o <div id="itens" class="mt-4"> no exemplo.html'
    novo_html = _BLOCO_ITENS.sub(lambda m: m.group(1) + tabela_estatica(ex, eol) + m.group(3), novo_html, count=1)
    faltam = [fr for fr in frases_que_a_pagina_tem_que_ter(ex, indiretos) if fr not in novo_html]
    if faltam:
        os.remove(tmp_xlsx)
        sys.exit("A prosa do exemplo.html não bate com os dados — atualize à mão e rode de novo:\n  "
                 + "\n  ".join(faltam))
    os.replace(tmp_xlsx, XLSX)
    tmp_html = PAGINA + ".tmp"
    with io.open(tmp_html, "w", encoding="utf-8", newline="") as f:
        f.write(novo_html)
    os.replace(tmp_html, PAGINA)
    print("exemplo regerado: %d itens (%d medidos), %d linhas de custo indireto"
          % (ex["total"], ex["medidos"], indiretos))


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--conferir", action="store_true", help="só confere; sai 1 se algo ficou velho")
    args = ap.parse_args()
    if args.conferir:
        probs = conferir()
        print("\n".join(probs) if probs else "exemplo em dia com o gerador")
        sys.exit(1 if probs else 0)
    regerar()
