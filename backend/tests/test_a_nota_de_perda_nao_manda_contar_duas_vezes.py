# -*- coding: utf-8 -*-
"""A nota de perdas da planilha não pode mandar somar a perda que a composição SINAPI já tem.

🩸 27/09/2026 — A revisão do plano de outubro do Instagram (achado PROD-14) achou a contradição:
o "Erro caro Nº 9" de 06/10 ensina que, no SINAPI, a perda de material já está nos coeficientes
das composições (Manual de Metodologias e Conceitos da Caixa) e que multiplicar por 1,10 "pra
garantir" conta a perda duas vezes. E a planilha para onde o post manda dizia, na nota 9:
"Perdas de material (5-10% típico) NÃO aplicadas automaticamente — adicionar ao preencher a
coluna de custo se pertinente" — ao lado de uma coluna REF que aponta justamente composição SINAPI.

O guarda cobra o COMPORTAMENTO da nota, não a frase: ela diz que a perda já está nos
coeficientes e não sugere percentual pra somar.
"""
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from models import BudgetItem, Confidence, ProjectData  # noqa: E402

# "5-10%", "5 a 10%", "5–10 %"
FAIXA_DE_PERCENTUAL = re.compile(r"\d+\s*(?:-|–|a)\s*\d+\s*%")


def _nota_de_perda_ok(texto):
    """None se a nota está certa; senão, o motivo."""
    if FAIXA_DE_PERCENTUAL.search(texto):
        return "sugere uma faixa de percentual de perda pra somar"
    if "coeficiente" not in texto.lower():
        return "não diz que a composição SINAPI já traz a perda nos coeficientes"
    return None


def _notas_de_perda(tmp_path):
    from openpyxl import load_workbook
    from spreadsheet import generate_spreadsheet
    caminho = str(tmp_path / "planilha.xlsx")
    itens = [BudgetItem(item_num="1.1", description="Piso porcelanato retificado 60x60", unit="m2",
                        quantity=96.3, discipline="Pisos e Rodapés", confidence=Confidence.ESTIMADO)]
    generate_spreadsheet(ProjectData(name="projeto de teste"), itens, caminho)
    return [str(c.value) for aba in load_workbook(caminho).worksheets for linha in aba.iter_rows()
            for c in linha if isinstance(c.value, str) and "perdas de material" in c.value.lower()]


def test_a_nota_de_perda_nao_manda_somar_a_perda_de_novo(tmp_path):
    notas = _notas_de_perda(tmp_path)
    assert notas, "a planilha saiu sem a nota de perdas — o guarda não achou o que conferir"
    problemas = ["%r: %s" % (n[:90], _nota_de_perda_ok(n)) for n in notas if _nota_de_perda_ok(n)]
    assert not problemas, "a nota de perdas voltou a contar a perda duas vezes: " + " | ".join(problemas)


def test_CONTROLE_o_guarda_REPROVA_a_nota_antiga():
    """🧪 Prova que morde, sem tocar em arquivo nenhum: a nota que estava no ar em 27/09."""
    velha = ("9. Perdas de material (5-10% típico) NÃO aplicadas automaticamente — adicionar ao "
             "preencher a coluna de custo se pertinente.")
    assert _nota_de_perda_ok(velha), "o guarda deixou passar a nota que estava no ar"
    assert _nota_de_perda_ok("9. Perdas de material: some 5 a 10 % ao custo."), "faixa com 'a'"
    assert _nota_de_perda_ok("9. Perdas de material NÃO aplicadas.") == (
        "não diz que a composição SINAPI já traz a perda nos coeficientes")
