# -*- coding: utf-8 -*-
"""A fase do cronograma cresce com a QUANTIDADE, não com o prazo — e o
coeficiente só conta o tipo de item que ele representa.

🩸 18/10/2026 — o /exemplo (reforma de escritório de 118,5 m², 5 meses) saía
com "Instalações hidráulicas: 82 dias" para 3 bacias + 1 cuba, "Divisórias e
vidros: 30 dias" para 8 portas e "Demolição: 22 dias" para 18 m² de divisória.
As três fases não tinham coeficiente pra unidade que o quantitativo traz e
caíam no % fixo do PRAZO ('padrão'): dobrar o prazo dobrava a fase.

Medido nos 23 cronogramas da base: 47 fases 'padrão' com quantidade contável
sem coeficiente, em 19 deles.

🪤 A 1ª versão do conserto deu coeficiente por UNIDADE e o acervo derrubou:
o "un" de hidráulica da base ia de luva e joelho (320 numa obra só) a estaca
raiz classificada na disciplina errada; o "un" de divisórias era quase todo
janela; "demolição m²" era terraplenagem (9.453 m² → 311 dias). Por isso o
coeficiente novo exige que a DESCRIÇÃO case com o item que ele representa.
"""
import os
import sys

_AQUI = os.path.dirname(os.path.abspath(__file__))
_BACKEND = os.path.dirname(_AQUI)
sys.path.insert(0, _BACKEND)

from cronograma import gerar_cronograma  # noqa: E402


def _it(disc, desc, unit, qty, selo="estimado"):
    return {"discipline": disc, "description": desc, "unit": unit,
            "quantity": qty, "confidence": selo}


def _fase(items, label, meses=5):
    c = gerar_cronograma(items, "2026-11-02", meses)
    return {f["label"]: f for f in c["fases"]}[label]


HIDR = "Instalações Hidráulicas"
DIVI = "Divisórias e Vidros"
DEMO = "Demolição e Remoção"

LOUCAS = [_it(HIDR, "Bacia sanitária com caixa acoplada", "un", 3),
          _it(HIDR, "Cuba de apoio + torneira (copa)", "un", 1)]


# ═══════════════════════════════════════════════════════════════════════════
#  O caso do /exemplo: as três fases saem da conta
# ═══════════════════════════════════════════════════════════════════════════

def test_LOUCA_da_hidraulica_entra_na_conta():
    f = _fase(LOUCAS, "Instalações hidráulicas")
    assert f["origem"] != "padrão", "3 bacias + 1 cuba seguem no % do prazo"
    assert "86888" in f["origem_detalhe"] and "4 un" in f["origem_detalhe"]
    assert f["dur_dias"] <= 7, f["dur_dias"]


def test_a_fase_NAO_cresce_com_o_prazo():
    """🩸 O defeito em uma linha: 5 meses davam 82 dias, 10 meses dariam 165."""
    assert (_fase(LOUCAS, "Instalações hidráulicas", 5)["dur_dias"]
            == _fase(LOUCAS, "Instalações hidráulicas", 10)["dur_dias"])


def test_PORTA_da_divisoria_entra_na_conta():
    f = _fase([_it(DIVI, "Porta de madeira 80x210 com batente", "un", 6, "confirmado"),
               _it(DIVI, "Porta de vidro temperado 90x210", "un", 2, "confirmado")],
              "Divisórias e vidros")
    assert f["origem"] == "calculada", f["origem"]
    assert "90822" in f["origem_detalhe"] and "8 un" in f["origem_detalhe"]


def test_DEMOLICAO_de_divisoria_entra_na_conta():
    f = _fase([_it(DEMO, "Remoção de divisória existente", "m²", 18)], "Demolição")
    assert f["origem"] == "estimada", f["origem"]
    assert "97638" in f["origem_detalhe"]
    assert f["dur_dias"] <= 7, f["dur_dias"]


# ═══════════════════════════════════════════════════════════════════════════
#  O que o acervo mostrou: item de outro tipo NÃO entra (fica como antes)
# ═══════════════════════════════════════════════════════════════════════════

def test_CONEXAO_nao_vira_louca():
    """🩸 320 luvas e joelhos × 1,443 Hh viravam 2 meses de hidráulica."""
    f = _fase([_it(HIDR, "Fornecimento e instalação de luva simples PVC", "un", 171),
               _it(HIDR, "Joelho 45° PVC esgoto ø50mm", "un", 149)],
              "Instalações hidráulicas")
    assert f["origem"] == "padrão", f.get("origem_detalhe")


def test_ESTACA_e_TANQUE_SEPTICO_nao_viram_louca():
    """Estaca raiz caiu na disciplina hidráulica num projeto real; tanque
    séptico não é tanque de lavar."""
    f = _fase([_it(HIDR, "Estaca raiz Ø20 cm, concreto fck 18 MPa", "un", 118),
               _it(HIDR, "Tanque séptico (várias variantes)", "un", 4)],
              "Instalações hidráulicas")
    assert f["origem"] == "padrão", f.get("origem_detalhe")


def test_PONTO_de_aparelho_conta_sem_acento_e_com_acento():
    f = _fase([_it(HIDR, "Ponto de lavatório — ramal de água fria", "un", 2),
               _it(HIDR, "PONTO DE MICTORIO", "un", 1)], "Instalações hidráulicas")
    assert "3 un" in f["origem_detalhe"], f["origem_detalhe"]


def test_JANELA_nao_vira_porta():
    f = _fase([_it(DIVI, "Janela maxim-ar — dimensões a confirmar", "un", 46)],
              "Divisórias e vidros")
    assert f["origem"] == "padrão", f.get("origem_detalhe")


def test_TERRAPLENAGEM_nao_vira_demolicao_de_parede():
    """🩸 9.453 m² de plataforma de terraplenagem → 311 dias de 'demolição'."""
    f = _fase([_it(DEMO, "Plataforma de terraplenagem — regularização", "m²", 9453),
               _it(DEMO, "Limpeza e preparo de terreno — remoção de vegetação", "m²", 4118)],
              "Demolição", meses=11)
    assert f["origem"] == "padrão", f.get("origem_detalhe")


# ═══════════════════════════════════════════════════════════════════════════
#  A fase que segue sem conta DIZ que não teve conta — na tela e no arquivo
# ═══════════════════════════════════════════════════════════════════════════

def test_fase_PADRAO_explica_que_nao_houve_conta():
    f = _fase([_it("Serviços Preliminares", "Administração da obra", "vb", 1)],
              "Serviços preliminares")
    assert f["origem"] == "padrão"
    assert "Sem quantidade" in (f.get("origem_detalhe") or ""), f.get("origem_detalhe")


def test_o_XLSX_leva_a_origem_de_cada_fase(tmp_path):
    from openpyxl import load_workbook
    from cronograma_xlsx import gerar_cronograma_xlsx

    c = gerar_cronograma(LOUCAS + [_it("Serviços Preliminares", "Mobilização", "vb", 1)],
                         "2026-11-02", 5)
    fases = c["fases"]
    for f in fases:
        f["valor_previsto"] = 1000.0
    c["financeiro"] = {"total_informado": 2000.0, "n_fases": 2, "n_fases_com_valor": 2}
    caminho = gerar_cronograma_xlsx(c, str(tmp_path / "c.xlsx"))
    ws = load_workbook(caminho).active
    linhas = [[x.value for x in r] for r in ws.iter_rows()]
    hdr = next(r for r in linhas if "ORIGEM DA DURAÇÃO" in r)
    i_or, i_pct = hdr.index("ORIGEM DA DURAÇÃO"), hdr.index("% EXECUTADO")
    i_val = hdr.index("VALOR DA FASE (R$)")
    corpo = {r[1]: r for r in linhas if r[1] in ("Instalações hidráulicas",
                                                  "Serviços preliminares")}
    assert corpo["Instalações hidráulicas"][i_or] == "calculada de estimativa"
    assert corpo["Serviços preliminares"][i_or].startswith("sem conta")
    # A coluna nova não pode empurrar o formato das vizinhas pro lugar errado.
    lin = next(r for r in ws.iter_rows() if r[1].value == "Instalações hidráulicas")
    assert lin[i_pct].number_format.endswith("%"), lin[i_pct].number_format
    assert "R$" in lin[i_val].number_format or "#,##0" in lin[i_val].number_format, \
        lin[i_val].number_format
