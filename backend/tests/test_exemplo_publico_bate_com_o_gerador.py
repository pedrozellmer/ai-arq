# -*- coding: utf-8 -*-
"""O exemplo público mostra a planilha que sai HOJE — e a página conta a mesma história que ela.

🩸 27/09/2026 — A planilha do exemplo era de 17/07. O gerador seguiu andando (24/08: ORIGEM DA
MEDIÇÃO e ESPECIFICAÇÃO, 9 → 11 colunas; 27/09: aba SINAPI sem "match", nota de perdas nova) e
ninguém regerou: quem baixava o exemplo pra "ver como volta" recebia uma planilha que não sai mais
(achado PROD-11 da revisão do plano de outubro do Instagram — o exemplo é o 1º link da bio o mês
inteiro). E página × planilha já tinham se descolado uma vez (d5b45d6, 17/07).

Agora as duas saem de UMA fonte — exemplos/itens.json — por scripts/gerar_exemplo_publico.py.
Mudou o gerador? Este guarda fica vermelho até alguém rodar:

    python scripts/gerar_exemplo_publico.py
"""
import importlib.util
import io
import json
import os
import re
import shutil

RAIZ = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
XLSX = os.path.join(RAIZ, "exemplos", "quantitativo-exemplo-aiarq.xlsx")
PAGINA = os.path.join(RAIZ, "exemplo.html")


def _script():
    spec = importlib.util.spec_from_file_location(
        "gerar_exemplo_publico", os.path.join(RAIZ, "scripts", "gerar_exemplo_publico.py"))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def _ex_da_pagina(caminho=PAGINA):
    html = io.open(caminho, encoding="utf-8").read()
    return json.loads(re.search(r"const EX = (\{.*?\]\});", html, re.S).group(1))


def _linhas_da_planilha():
    """{nº: (descrição, un, qtde, observações, origem da medição, ref)} da aba Orçamento."""
    from openpyxl import load_workbook
    ws = load_workbook(XLSX)["Orçamento"]
    cab = next(r for r in ws.iter_rows(values_only=True) if r[0] == "ITEM")
    col = {nome: i for i, nome in enumerate(cab)}
    linhas = {}
    for r in ws.iter_rows(values_only=True):
        # 0.x são as PREMISSAS (área, "também gera…") — metadado, não item
        if r[0] and re.match(r"^[1-9]\d*\.\d+$", str(r[0])):
            linhas[str(r[0])] = (r[col["DESCRIÇÃO DO SERVIÇO"]], r[col["UN"]], r[col["QTDE"]],
                                 r[col["OBSERVAÇÕES"]], r[col["ORIGEM DA MEDIÇÃO"]], r[col["REF."]])
    return linhas


def test_o_exemplo_esta_em_dia_com_o_gerador():
    problemas = _script().conferir()
    assert not problemas, (
        "o exemplo público ficou velho em relação ao gerador de planilha — rode "
        "`python scripts/gerar_exemplo_publico.py` e suba a planilha e o exemplo.html juntos:\n  "
        + "\n  ".join(problemas))


def test_a_planilha_do_exemplo_tem_as_colunas_de_hoje():
    from openpyxl import load_workbook
    ws = load_workbook(XLSX)["Orçamento"]
    cab = next(r for r in ws.iter_rows(values_only=True) if r[0] == "ITEM")
    for coluna in ("ORIGEM DA MEDIÇÃO", "ESPECIFICAÇÃO", "OBSERVAÇÕES", "REF."):
        assert coluna in cab, "a planilha do exemplo não tem a coluna %r: %r" % (coluna, cab)


def test_pagina_e_planilha_contam_a_mesma_historia():
    """Conferência independente do script: lê os DOIS arquivos publicados, item a item."""
    ex, linhas = _ex_da_pagina(), _linhas_da_planilha()
    assert len(ex["itens"]) == ex["total"] == len(linhas), (len(ex["itens"]), ex["total"], len(linhas))
    for it in ex["itens"]:
        assert it["num"] in linhas, "a página mostra o item %s, que a planilha não tem" % it["num"]
        desc, un, qtd, obs, _origem, ref = linhas[it["num"]]
        assert (desc, un, float(qtd)) == (it["desc"], it["un"], float(it["qtd"])), (it["num"], desc, it["desc"])
        selo = "✓ MEDIDO" if it["medido"] else "⚠ ESTIMADO"
        assert str(obs).startswith(selo), "%s: a página diz %s e a planilha diz %r" % (it["num"], selo, obs[:30])
        if it["sinapi"]:
            assert it["sinapi"]["cod"] in str(ref), "%s: a página mostra SINAPI %s e a REF diz %r" % (
                it["num"], it["sinapi"]["cod"], ref)
    assert ex["medidos"] == sum(1 for i in ex["itens"] if i["medido"])


def test_todo_medido_do_exemplo_diz_de_onde_veio():
    """Regra dura nº1 na vitrine: selo de MEDIDO sem procedência é o que a planilha avisa como
    '⚠ sem procedência registrada'. O exemplo não pode ensinar isso como normal."""
    ruins = ["%s %r" % (num, origem) for num, (_d, _u, _q, obs, origem, _r) in _linhas_da_planilha().items()
             if str(obs).startswith("✓") and (not origem or "sem procedência" in str(origem) or origem == "—")]
    assert not ruins, "item MEDIDO do exemplo sem procedência: %s" % ruins


def test_CONTROLE_o_conferir_pega_planilha_e_pagina_mexidas(tmp_path):
    """🧪 Prova que morde: uma quantidade trocada na planilha e um número trocado na página."""
    from openpyxl import load_workbook
    mod = _script()
    xlsx = str(tmp_path / "exemplo.xlsx")
    shutil.copy(XLSX, xlsx)
    wb = load_workbook(xlsx)
    ws = wb["Orçamento"]
    alvo = next(c for linha in ws.iter_rows() for c in linha[3:4]
                if isinstance(c.value, (int, float)) and linha[0].value and re.match(r"^\d+\.\d+$", str(linha[0].value)))
    alvo.value = alvo.value + 1
    wb.save(xlsx)
    assert any(p.startswith("planilha:") for p in mod.conferir(xlsx=xlsx)), "o conferir não viu a quantidade mexida"

    pagina = str(tmp_path / "exemplo.html")
    html = io.open(PAGINA, encoding="utf-8").read()
    io.open(pagina, "w", encoding="utf-8").write(html.replace('"num": "1.1"', '"num": "1.9"', 1))
    assert any(p.startswith("página:") for p in mod.conferir(pagina=pagina)), "o conferir não viu a página mexida"
